#!/usr/bin/env python3
"""Extract a lecture to plain text once, so it is read from text instead of re-opening the source.

Usage: python extract_lecture.py <file_or_folder> <out_dir>
Writes <out_dir>/lecture.txt with "=== page N ===" markers (or "=== slide N ===")
and <out_dir>/lecture_meta.json listing pages with little text (likely diagrams or
scanned images) that should be looked at as images. Handles .pdf .pptx .docx .md .txt .tex .ipynb.
A folder is treated as one lecture (files in name order).
"""
import sys, json, pathlib, subprocess

def pdf(p):
    try:
        import pdfplumber
        with pdfplumber.open(p) as d:
            return [(pg.extract_text() or "") for pg in d.pages]
    except Exception:
        out = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True).stdout
        return out.split("\f")

def pptx(p):
    from pptx import Presentation
    pages = []
    for s in Presentation(p).slides:
        t = []
        for sh in s.shapes:
            if sh.has_text_frame: t.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows: t.append(" | ".join(c.text for c in r.cells))
        if s.has_notes_slide: t.append("[notes] " + s.notes_slide.notes_text_frame.text)
        pages.append("\n".join(t))
    return pages

def docx(p):
    import docx as d
    return ["\n".join(x.text for x in d.Document(p).paragraphs)]

def ipynb(p):
    nb = json.loads(pathlib.Path(p).read_text())
    return ["".join(c.get("source", "")) for c in nb.get("cells", [])]

def text(p): return [pathlib.Path(p).read_text(errors="ignore")]

H = {".pdf": (pdf, "page"), ".pptx": (pptx, "slide"), ".docx": (docx, "part"), ".ipynb": (ipynb, "cell"),
     ".md": (text, "part"), ".txt": (text, "part"), ".tex": (text, "part")}

src = pathlib.Path(sys.argv[1]); out = pathlib.Path(sys.argv[2]); out.mkdir(parents=True, exist_ok=True)
files = sorted(f for f in (src.rglob("*") if src.is_dir() else [src]) if f.suffix.lower() in H)
buf, sparse, total = [], [], 0
for f in files:
    fn, unit = H[f.suffix.lower()]
    try: pages = fn(f)
    except Exception as e:
        buf.append(f"### {f.name}: could not read ({e})"); continue
    buf.append(f"##### FILE {f.name}")
    for i, t in enumerate(pages, 1):
        total += 1
        buf.append(f"=== {unit} {i} ===\n{t.strip()}")
        if f.suffix.lower() in (".pdf", ".pptx") and len(t.strip()) < 80: sparse.append(f"{f.name} {unit} {i}")
(out/"lecture.txt").write_text("\n\n".join(buf))
(out/"lecture_meta.json").write_text(json.dumps({"files": [f.name for f in files], "units": total,
    "look_at_as_images": sparse}, indent=1))
print(f"{len(files)} file(s), {total} pages/slides -> {out/'lecture.txt'}")
if sparse: print(f"{len(sparse)} with little text (view these as images): " + ", ".join(sparse[:20]) + (" ..." if len(sparse) > 20 else ""))
